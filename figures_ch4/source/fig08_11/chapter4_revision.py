from __future__ import annotations

import argparse
import csv
import json
import math
import re
from copy import deepcopy
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
CODE_ROOT = next(ROOT.glob("*/CWT-TSI-new"))
WORK = ROOT / "work" / "chapter4_revision"
EVIDENCE_DIR = WORK / "evidence"
FIGURE_DIR = WORK / "figures"
SOURCE_DOCX = ROOT / "CWT修改版中文稿.docx"
OUTPUT_DOCX = ROOT / "CWT修改版中文稿_第四章期刊化修订稿.docx"

COMPARISON_CSV = CODE_ROOT / "log" / "comparison_table4_5_full" / "comparison_table.csv"
ABLATION_CSV = (
    CODE_ROOT
    / "log"
    / "ablation_table7_full_20260706_0229"
    / "ablation_table7.csv"
)
COMPARISON_FIGURE_DIR = (
    CODE_ROOT / "log" / "comparison_table4_5_full" / "journal_figures"
)
ABLATION_FIGURE_DIR = (
    CODE_ROOT
    / "log"
    / "ablation_table7_full_20260706_0229"
    / "journal_figures"
)

DATASETS = ["XMBRT", "HZMetro", "BJMetro"]
MODELS = [
    "LSTM",
    "ASTGCN-r",
    "GraphWaveNet",
    "GMAN",
    "STSGCN",
    "StemGNN",
    "STFGNN",
    "ASTGNN",
    "CWT-TSI",
]
VARIANTS = ["No-Img", "No-TS", "No-Fuse", "CWT-TSI"]
METRICS = ["mae", "rmse", "mape"]
METRIC_LABELS = {"mae": "MAE", "rmse": "RMSE", "mape": "MAPE (%)"}
RUNS = {
    "XMBRT": "xmbrt_full",
    "HZMetro": "hzmetro_full",
    "BJMetro": "bjmetro_full",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def metric_triplet(target: np.ndarray, prediction: np.ndarray) -> dict[str, float]:
    target = np.asarray(target, dtype=np.float64).copy()
    prediction = np.asarray(prediction, dtype=np.float64)
    target[np.abs(target) < 1e-4] = 0.0
    diff = prediction - target
    mae = float(np.mean(np.abs(diff)))
    rmse = float(np.sqrt(np.mean(np.square(diff))))
    valid = target != 0
    if not np.any(valid):
        mape = 0.0
    else:
        mape = float(np.mean(np.abs(diff[valid] / target[valid])) * 100.0)
    return {"mae": mae, "rmse": rmse, "mape": mape}


def assert_close(actual: float, expected: float, label: str, atol: float = 5e-4) -> None:
    if not math.isclose(actual, expected, rel_tol=0.0, abs_tol=atol):
        raise AssertionError(f"{label}: actual={actual}, expected={expected}")


def build_evidence() -> dict:
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    comparison_rows = read_csv(COMPARISON_CSV)
    ablation_rows = read_csv(ABLATION_CSV)
    if len(comparison_rows) != 27:
        raise AssertionError(f"comparison rows: {len(comparison_rows)} != 27")
    if len(ablation_rows) != 12:
        raise AssertionError(f"ablation rows: {len(ablation_rows)} != 12")

    comparison: dict[str, dict[str, dict[str, float]]] = {d: {} for d in DATASETS}
    for row in comparison_rows:
        dataset = row["dataset"]
        model = row["model"]
        comparison[dataset][model] = {m: float(row[m]) for m in METRICS}
    for dataset in DATASETS:
        if list(comparison[dataset].keys()) != MODELS:
            raise AssertionError(f"{dataset}: unexpected comparison model order")

    ablation: dict[str, dict[str, dict[str, float]]] = {d: {} for d in DATASETS}
    for row in ablation_rows:
        dataset = row["dataset"]
        model = row["model"]
        ablation[dataset][model] = {m: float(row[m]) for m in METRICS}
    for dataset in DATASETS:
        if set(ablation[dataset]) != set(VARIANTS):
            raise AssertionError(f"{dataset}: unexpected ablation variants")

    full_runs: dict[str, dict] = {}
    for dataset, run in RUNS.items():
        run_dir = CODE_ROOT / "save" / run
        metrics = json.loads((run_dir / "metrics.json").read_text(encoding="utf-8"))
        if len(metrics.get("per_step_metrics", [])) != 12:
            raise AssertionError(f"{run}: per_step_metrics does not contain 12 rows")
        prediction = np.load(run_dir / "test_predict.npy")
        target = np.load(run_dir / "test_target.npy")
        if prediction.shape != target.shape:
            raise AssertionError(f"{run}: prediction/target shapes differ")
        if prediction.ndim != 2 or prediction.shape[1] != 12:
            raise AssertionError(f"{run}: unexpected prediction shape {prediction.shape}")

        recomputed = metric_triplet(target, prediction)
        for metric in METRICS:
            assert_close(
                recomputed[metric],
                float(metrics[metric]),
                f"{run} overall {metric}",
            )

        recomputed_per_step = []
        for step in range(12):
            step_metrics = metric_triplet(target[:, step], prediction[:, step])
            expected = metrics["per_step_metrics"][step]
            for metric in METRICS:
                assert_close(
                    step_metrics[metric],
                    float(expected[metric]),
                    f"{run} step {step + 1} {metric}",
                )
            recomputed_per_step.append({"step": step + 1, **step_metrics})

        history_path = run_dir / "history.npz"
        history = np.load(history_path)
        train_loss = history["train_loss"].astype(float).tolist()
        val_loss = history["val_loss"].astype(float).tolist()
        if len(train_loss) != len(val_loss) or len(train_loss) == 0:
            raise AssertionError(f"{run}: invalid loss history")

        full_runs[dataset] = {
            "run": run,
            "metrics": {m: float(metrics[m]) for m in METRICS},
            "per_step_metrics": [
                {
                    "step": int(row["step"]),
                    **{m: float(row[m]) for m in METRICS},
                }
                for row in metrics["per_step_metrics"]
            ],
            "recomputed_per_step_metrics": recomputed_per_step,
            "train_loss": train_loss,
            "val_loss": val_loss,
            "environment": metrics["environment"],
            "history_step": int(metrics["history_step"]),
            "target_step": int(metrics["target_step"]),
            "num_nodes": int(metrics["num_nodes"]),
            "num_epochs": int(metrics["num_epochs"]),
            "batch_size": int(metrics["batch_size"]),
            "learning_rate": float(metrics["learning_rate"]),
        }

    comparison_relative = []
    ranks: dict[str, dict[str, dict[str, int]]] = {d: {} for d in DATASETS}
    for dataset in DATASETS:
        for metric in METRICS:
            ordered = sorted(MODELS, key=lambda model: comparison[dataset][model][metric])
            for rank, model in enumerate(ordered, start=1):
                ranks[dataset].setdefault(model, {})[metric] = rank
            baselines = [m for m in MODELS if m != "CWT-TSI"]
            best_baseline = min(baselines, key=lambda model: comparison[dataset][model][metric])
            best_value = comparison[dataset][best_baseline][metric]
            cwt_value = comparison[dataset]["CWT-TSI"][metric]
            improvement = (best_value - cwt_value) / best_value * 100.0
            comparison_relative.append(
                {
                    "dataset": dataset,
                    "metric": METRIC_LABELS[metric],
                    "best_baseline": best_baseline,
                    "best_value": best_value,
                    "cwt_value": cwt_value,
                    "improvement": improvement,
                    "cwt_rank": ranks[dataset]["CWT-TSI"][metric],
                }
            )

    with (EVIDENCE_DIR / "comparison_relative.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=list(comparison_relative[0].keys()))
        writer.writeheader()
        writer.writerows(comparison_relative)

    ablation_relative = []
    for dataset in DATASETS:
        full = ablation[dataset]["CWT-TSI"]
        for variant in ["No-Img", "No-TS", "No-Fuse"]:
            for metric in METRICS:
                change = (ablation[dataset][variant][metric] - full[metric]) / full[metric] * 100.0
                ablation_relative.append(
                    {
                        "dataset": dataset,
                        "variant": variant,
                        "metric": METRIC_LABELS[metric],
                        "full_value": full[metric],
                        "variant_value": ablation[dataset][variant][metric],
                        "relative_change": change,
                    }
                )

    with (EVIDENCE_DIR / "ablation_relative.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=list(ablation_relative[0].keys()))
        writer.writeheader()
        writer.writerows(ablation_relative)

    evidence = {
        "comparison": comparison,
        "comparison_ranks": ranks,
        "comparison_relative": comparison_relative,
        "ablation": ablation,
        "ablation_relative": ablation_relative,
        "full_runs": full_runs,
    }
    (EVIDENCE_DIR / "chapter4_evidence.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        "comparison_rows=27 ablation_rows=12 metric_recompute=PASS "
        f"evidence={EVIDENCE_DIR / 'chapter4_evidence.json'}"
    )
    return evidence


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    names = ["arialbd.ttf", "Arial Bold.ttf"] if bold else ["arial.ttf", "Arial.ttf"]
    for name in names:
        path = Path("C:/Windows/Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default(size=size)


def dashed_horizontal(
    draw: ImageDraw.ImageDraw,
    x0: float,
    x1: float,
    y: float,
    fill: str,
    width: int = 2,
    dash: int = 10,
) -> None:
    x = x0
    while x < x1:
        draw.line((x, y, min(x + dash, x1), y), fill=fill, width=width)
        x += dash * 1.8


def draw_line_chart(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    xs: list[float],
    series: list[tuple[str, list[float], str]],
    title: str,
    x_label: str,
    y_label: str,
    average: float | None = None,
    x_ticks: list[int] | None = None,
) -> None:
    left, top, right, bottom = box
    title_font = load_font(27, bold=True)
    label_font = load_font(21)
    tick_font = load_font(18)
    small_font = load_font(17)
    plot_left = left + 105
    plot_right = right - 30
    plot_top = top + 64
    plot_bottom = bottom - 78
    values = [value for _, data, _ in series for value in data]
    if average is not None:
        values.append(average)
    y_min = min(values)
    y_max = max(values)
    span = y_max - y_min
    if span <= 1e-12:
        span = max(abs(y_max), 1.0)
    pad = span * 0.12
    y_min -= pad
    y_max += pad
    x_min, x_max = min(xs), max(xs)

    def px(value: float) -> float:
        return plot_left + (value - x_min) / (x_max - x_min) * (plot_right - plot_left)

    def py(value: float) -> float:
        return plot_bottom - (value - y_min) / (y_max - y_min) * (plot_bottom - plot_top)

    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(
        ((left + right - (title_box[2] - title_box[0])) / 2, top + 8),
        title,
        fill="#252525",
        font=title_font,
    )
    draw.line((plot_left, plot_top, plot_left, plot_bottom), fill="#777777", width=2)
    draw.line((plot_left, plot_bottom, plot_right, plot_bottom), fill="#777777", width=2)
    for idx in range(5):
        value = y_min + idx * (y_max - y_min) / 4
        y = py(value)
        dashed_horizontal(draw, plot_left, plot_right, y, "#DDDDDD", width=1, dash=8)
        label = f"{value:.2f}" if abs(value) < 100 else f"{value:.1f}"
        bbox = draw.textbbox((0, 0), label, font=tick_font)
        draw.text((plot_left - 12 - (bbox[2] - bbox[0]), y - 10), label, fill="#555555", font=tick_font)
    ticks = x_ticks or [int(x_min), int(round((x_min + x_max) / 2)), int(x_max)]
    for tick in ticks:
        x = px(tick)
        draw.line((x, plot_bottom, x, plot_bottom + 6), fill="#777777", width=2)
        label = str(tick)
        bbox = draw.textbbox((0, 0), label, font=tick_font)
        draw.text((x - (bbox[2] - bbox[0]) / 2, plot_bottom + 10), label, fill="#555555", font=tick_font)
    for label, data, color in series:
        points = [(px(x), py(y)) for x, y in zip(xs, data)]
        draw.line(points, fill=color, width=4, joint="curve")
        if len(points) <= 12:
            for point in points:
                draw.ellipse(
                    (point[0] - 5, point[1] - 5, point[0] + 5, point[1] + 5),
                    fill=color,
                    outline="white",
                    width=1,
                )
    if average is not None:
        y = py(average)
        dashed_horizontal(draw, plot_left, plot_right, y, "#555555", width=2, dash=12)
        draw.text((plot_right - 145, y - 24), "12-step avg.", fill="#555555", font=small_font)

    x_box = draw.textbbox((0, 0), x_label, font=label_font)
    draw.text(
        ((plot_left + plot_right - (x_box[2] - x_box[0])) / 2, bottom - 30),
        x_label,
        fill="#333333",
        font=label_font,
    )
    draw.text((left + 8, plot_top - 34), y_label, fill="#333333", font=label_font)
    if len(series) > 1:
        legend_x = plot_left + 12
        legend_y = plot_top + 8
        for label, _, color in series:
            draw.line((legend_x, legend_y + 9, legend_x + 34, legend_y + 9), fill=color, width=4)
            draw.text((legend_x + 43, legend_y), label, fill="#333333", font=small_font)
            legend_y += 28


def generate_figures(evidence: dict) -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    colors = {"XMBRT": "#2F6B9A", "HZMetro": "#E07A2D", "BJMetro": "#3D8B6D"}

    width, height = 3600, 3000
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    main_font = load_font(48, bold=True)
    main_title = "CWT-TSI Performance Across Prediction Horizons"
    bbox = draw.textbbox((0, 0), main_title, font=main_font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 24), main_title, fill="#222222", font=main_font)
    top_offset = 95
    cell_w = 1160
    cell_h = 930
    for row_idx, dataset in enumerate(DATASETS):
        rows = evidence["full_runs"][dataset]["per_step_metrics"]
        steps = [row["step"] for row in rows]
        for col_idx, metric in enumerate(METRICS):
            left = 45 + col_idx * cell_w
            top = top_offset + row_idx * cell_h
            box = (left, top, left + cell_w - 35, top + cell_h - 25)
            values = [row[metric] for row in rows]
            draw_line_chart(
                draw,
                box,
                steps,
                [(dataset, values, colors[dataset])],
                METRIC_LABELS[metric],
                "Prediction step",
                dataset,
                average=evidence["full_runs"][dataset]["metrics"][metric],
                x_ticks=[1, 3, 6, 9, 12],
            )
    per_step_path = FIGURE_DIR / "cwt_per_step_metrics.png"
    image.save(per_step_path, dpi=(240, 240))

    width, height = 3600, 1220
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    main_title = "Training and Validation Loss of CWT-TSI"
    bbox = draw.textbbox((0, 0), main_title, font=main_font)
    draw.text(((width - (bbox[2] - bbox[0])) / 2, 24), main_title, fill="#222222", font=main_font)
    cell_w = 1160
    for col_idx, dataset in enumerate(DATASETS):
        train_loss = evidence["full_runs"][dataset]["train_loss"]
        val_loss = evidence["full_runs"][dataset]["val_loss"]
        epochs = list(range(1, len(train_loss) + 1))
        left = 45 + col_idx * cell_w
        box = (left, 100, left + cell_w - 35, 1170)
        draw_line_chart(
            draw,
            box,
            epochs,
            [
                ("Training loss", train_loss, "#2F6B9A"),
                ("Validation loss", val_loss, "#C4472D"),
            ],
            dataset,
            "Epoch",
            "Huber loss",
            x_ticks=[1, max(2, len(epochs) // 2), len(epochs)],
        )
    convergence_path = FIGURE_DIR / "cwt_training_convergence.png"
    image.save(convergence_path, dpi=(240, 240))

    for path in [per_step_path, convergence_path]:
        with Image.open(path) as check_image:
            if check_image.width < 2400 or check_image.mode != "RGB":
                raise AssertionError(
                    f"invalid generated figure: {path} {check_image.size} {check_image.mode}"
                )
    print(f"figures=PASS per_step={per_step_path} convergence={convergence_path}")


def remove_paragraph_content(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_text(cell, text: str, size: float = 8.0, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    remove_paragraph_content(paragraph)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph: Paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def body_paragraph(doc: Document, anchor: Paragraph, text: str) -> Paragraph:
    paragraph = doc.add_paragraph(style="Normal")
    anchor._p.addprevious(paragraph._p)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def heading_paragraph(doc: Document, anchor: Paragraph, text: str, level: int) -> Paragraph:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    anchor._p.addprevious(paragraph._p)
    run = paragraph.add_run(text)
    run.font.name = "黑体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    return paragraph


def caption_paragraph(
    doc: Document, anchor: Paragraph, text: str, is_table: bool = False
) -> Paragraph:
    paragraph = doc.add_paragraph(style="Normal")
    anchor._p.addprevious(paragraph._p)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(3 if not is_table else 6)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph, is_table)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.0)
    return paragraph


def add_figure(
    doc: Document, anchor: Paragraph, image_path: Path, width_cm: float = 15.5
) -> Paragraph:
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph(style="Normal")
    anchor._p.addprevious(paragraph._p)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    set_keep_with_next(paragraph, True)
    return paragraph


def insert_table_clone(doc: Document, anchor: Paragraph, template_xml) -> Table:
    table_xml = deepcopy(template_xml)
    anchor._p.addprevious(table_xml)
    table = Table(table_xml, doc)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    return table


def format_table(table: Table, header_rows: int = 1, font_size: float = 8.0) -> None:
    for row_idx, row in enumerate(table.rows):
        if row_idx < header_rows:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(row_idx < header_rows))
        tr_pr = row._tr.get_or_add_trPr()
        for height in tr_pr.findall(qn("w:trHeight")):
            tr_pr.remove(height)


def mark_best_and_second(table: Table, data_start_row: int, values_by_col: dict[int, list[float]]) -> None:
    for col_idx, values in values_by_col.items():
        order = sorted(range(len(values)), key=lambda i: values[i])
        for rank, offset in enumerate(order[:2], start=1):
            cell = table.cell(data_start_row + offset, col_idx)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if rank == 1:
                        run.bold = True
                    else:
                        run.underline = True


def write_table1(table: Table) -> None:
    rows = [
        ["数据集", "城市/场景", "节点数", "总时间步", "有效样本数", "时间分辨率", "每日时间片"],
        ["XMBRT", "厦门BRT", "44", "4,320", "4,297", "4分钟", "216"],
        ["HZMetro", "杭州地铁", "80", "5,616", "5,593", "4分钟", "216"],
        ["BJMetro", "北京地铁", "276", "2,700", "2,677", "5分钟", "108"],
    ]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            set_cell_text(table.cell(row_idx, col_idx), value, size=8.5, bold=(row_idx == 0))
    format_table(table, header_rows=1, font_size=8.5)


def write_table2(table: Table) -> None:
    rows = [
        ["参数", "设定值"],
        ["硬件", "NVIDIA GeForce RTX 5060 Laptop GPU"],
        ["软件环境", "Python 3.10.20；PyTorch 2.11.0+cu128；CUDA 12.8"],
        ["优化器", "AdamW"],
        ["初始学习率", "0.001"],
        ["学习率调度", "CosineAnnealingLR（T_max=10）"],
        ["损失函数", "HuberLoss"],
        ["批大小", "16"],
        ["最大训练轮次", "100"],
        ["早停耐心值", "20"],
        ["历史/预测步长", "12/12"],
        ["CWT设置", "Morlet（morl）；尺度1–128；flatten模式"],
        ["图像尺寸", "1×128×128"],
        ["随机种子", "1"],
        ["数据划分", "训练集/验证集/测试集=6:2:2（按时间顺序）"],
    ]
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            set_cell_text(table.cell(row_idx, col_idx), value, size=9.0, bold=(row_idx == 0))
    format_table(table, header_rows=1, font_size=9.0)


def write_comparison_table(table: Table, comparison: dict) -> None:
    for model_idx, model in enumerate(MODELS, start=2):
        set_cell_text(table.cell(model_idx, 0), model, size=7.5, bold=(model == "CWT-TSI"))
        col = 1
        for dataset in DATASETS:
            for metric in METRICS:
                value = comparison[dataset][model][metric]
                set_cell_text(table.cell(model_idx, col), f"{value:.4f}", size=7.2)
                col += 1
    format_table(table, header_rows=2, font_size=7.2)
    values_by_col = {}
    col = 1
    for dataset in DATASETS:
        for metric in METRICS:
            values_by_col[col] = [comparison[dataset][model][metric] for model in MODELS]
            col += 1
    mark_best_and_second(table, 2, values_by_col)


def write_step_table(table: Table, run: dict) -> None:
    header = ["预测步长", "MAE", "RMSE", "MAPE (%)"]
    for col_idx, value in enumerate(header):
        set_cell_text(table.cell(0, col_idx), value, size=8.5, bold=True)
    for row_idx, row in enumerate(run["per_step_metrics"], start=1):
        values = [str(row["step"]), f"{row['mae']:.4f}", f"{row['rmse']:.4f}", f"{row['mape']:.4f}"]
        for col_idx, value in enumerate(values):
            set_cell_text(table.cell(row_idx, col_idx), value, size=8.3)
    avg = run["metrics"]
    values = ["平均值", f"{avg['mae']:.4f}", f"{avg['rmse']:.4f}", f"{avg['mape']:.4f}"]
    for col_idx, value in enumerate(values):
        set_cell_text(table.cell(13, col_idx), value, size=8.3, bold=True)
    format_table(table, header_rows=1, font_size=8.3)


def write_ablation_table(table: Table, ablation: dict) -> None:
    header = ["数据集", "模型", "MAE", "RMSE", "MAPE (%)"]
    for col_idx, value in enumerate(header):
        set_cell_text(table.cell(0, col_idx), value, size=8.5, bold=True)
    row_idx = 1
    values_by_dataset = {}
    for dataset in DATASETS:
        values_by_dataset[dataset] = {metric: [] for metric in METRICS}
        for variant in VARIANTS:
            values = [
                dataset,
                variant,
                f"{ablation[dataset][variant]['mae']:.4f}",
                f"{ablation[dataset][variant]['rmse']:.4f}",
                f"{ablation[dataset][variant]['mape']:.4f}",
            ]
            for col_idx, value in enumerate(values):
                set_cell_text(
                    table.cell(row_idx, col_idx),
                    value,
                    size=8.2,
                    bold=(variant == "CWT-TSI" and col_idx == 1),
                )
            for metric in METRICS:
                values_by_dataset[dataset][metric].append(ablation[dataset][variant][metric])
            row_idx += 1
    format_table(table, header_rows=1, font_size=8.2)
    start = 1
    for dataset in DATASETS:
        columns = {
            2: values_by_dataset[dataset]["mae"],
            3: values_by_dataset[dataset]["rmse"],
            4: values_by_dataset[dataset]["mape"],
        }
        mark_best_and_second(table, start, columns)
        start += 4


def relative_lookup(evidence: dict, dataset: str, metric_label: str) -> dict:
    return next(
        row
        for row in evidence["comparison_relative"]
        if row["dataset"] == dataset and row["metric"] == metric_label
    )


def ablation_change(evidence: dict, dataset: str, variant: str, metric_label: str) -> float:
    return next(
        row["relative_change"]
        for row in evidence["ablation_relative"]
        if row["dataset"] == dataset
        and row["variant"] == variant
        and row["metric"] == metric_label
    )


def build_docx(evidence: dict) -> Path:
    doc = Document(str(SOURCE_DOCX))
    chapter4 = next(p for p in doc.paragraphs if p.text.strip().startswith("4 实验结果与分析"))
    chapter5 = next(p for p in doc.paragraphs if p.text.strip().startswith("5 结论与展望"))
    formula_xml = [
        deepcopy(p._p)
        for idx, p in enumerate(doc.paragraphs)
        if p.style.name == "公式" and 112 <= idx < 176
    ]
    if len(formula_xml) != 3:
        raise AssertionError(f"expected 3 formula paragraphs, found {len(formula_xml)}")
    table_templates = [deepcopy(table._tbl) for table in doc.tables]
    if len(table_templates) != 7:
        raise AssertionError(f"expected 7 source tables, found {len(table_templates)}")

    node = chapter4._p.getnext()
    while node is not None and node is not chapter5._p:
        next_node = node.getnext()
        doc.element.body.remove(node)
        node = next_node

    heading_paragraph(doc, chapter5, "4.1  实验数据集", 2)
    body_paragraph(
        doc,
        chapter5,
        "本文在XMBRT、HZMetro和BJMetro三个真实公共交通客流数据集上开展实验。三个数据集分别对应厦门快速公交、杭州地铁和北京地铁，节点规模由44扩展至276，既覆盖不同交通方式，也覆盖小、中、大规模网络。该设置用于检验模型在不同客流量级与网络复杂度下的适用性。所有样本均按时间顺序以6:2:2划分为训练集、验证集和测试集，划分过程中不进行跨时间随机打乱，以避免未来信息泄漏。",
    )
    heading_paragraph(doc, chapter5, "4.1.1  XMBRT数据集", 3)
    body_paragraph(
        doc,
        chapter5,
        "XMBRT数据集包含厦门快速公交系统44个站点的客流记录，时间分辨率为4 min。原始序列共包含4,320个时间步；在构造12步历史输入和12步预测目标后得到4,297个有效样本。该数据集节点规模较小，客流具有较明显的日内周期，适合考察模型在小规模公交网络中的预测表现。",
    )
    heading_paragraph(doc, chapter5, "4.1.2  HZMetro数据集", 3)
    body_paragraph(
        doc,
        chapter5,
        "HZMetro数据集覆盖杭州地铁80个站点，时间分辨率为4 min，共包含5,616个时间步和5,593个有效样本。与XMBRT相比，该数据集具有更复杂的换乘关系和更大的节点规模，可用于评估模型对中等规模轨道交通网络客流模式的建模能力。",
    )
    heading_paragraph(doc, chapter5, "4.1.3  BJMetro数据集", 3)
    body_paragraph(
        doc,
        chapter5,
        "BJMetro数据集包含北京地铁276个站点，时间分辨率为5 min，共包含2,700个时间步和2,677个有效样本。该数据集节点数显著高于另外两个数据集，可用于检验模型在大规模交通网络和高客流量级条件下的性能。三个数据集的统计信息见表1。",
    )
    caption_paragraph(doc, chapter5, "表1  实验数据集基本统计信息", is_table=True)
    table1 = insert_table_clone(doc, chapter5, table_templates[0])
    write_table1(table1)

    heading_paragraph(doc, chapter5, "4.2  对比基线模型", 2)
    body_paragraph(
        doc,
        chapter5,
        "为从不同技术路线评估CWT-TSI，本文选取8种基线模型。LSTM[11]作为经典循环神经网络基线，仅基于时间序列进行预测；ASTGCN-r通过空间注意力、时间注意力与图卷积联合建模交通状态；Graph WaveNet[26]利用自适应邻接矩阵和扩张因果卷积学习时空依赖；GMAN[27]采用多头时空注意力及变换注意力完成多步预测。",
    )
    body_paragraph(
        doc,
        chapter5,
        "其余图模型包括STSGCN[23]、StemGNN[6]、STFGNN[16]和ASTGNN[10]。STSGCN通过局部时空图同步聚合邻近时间步与节点信息；StemGNN在谱域联合建模多元序列；STFGNN通过时空融合图结合空间邻接与时间相似性；ASTGNN采用注意力机制学习动态时空关系。对比实验中的所有模型均使用相同的数据划分、历史步长、预测步长和评价指标，结果取各自验证损失最优检查点在测试集上的一次评估值。",
    )

    heading_paragraph(doc, chapter5, "4.3  评价指标", 2)
    body_paragraph(
        doc,
        chapter5,
        "本文采用平均绝对误差（MAE）、均方根误差（RMSE）和平均绝对百分比误差（MAPE）评价预测性能。三项指标的定义分别为：",
    )
    for xml in formula_xml:
        chapter5._p.addprevious(deepcopy(xml))
    body_paragraph(
        doc,
        chapter5,
        "其中，ŷ_i和y_i分别表示第i个预测值和真实值，n表示参与计算的有效观测数。MAE反映误差的平均绝对量级；RMSE对较大误差赋予更高权重；MAPE衡量相对误差，计算时按照实验代码的掩码规则排除真实值为0的观测。三项指标均为越小越优。",
    )

    heading_paragraph(doc, chapter5, "4.4  实验环境与参数设置", 2)
    body_paragraph(
        doc,
        chapter5,
        "实验在NVIDIA GeForce RTX 5060 Laptop GPU上完成，软件环境为Python 3.10.20、PyTorch 2.11.0+cu128和CUDA 12.8。CWT-TSI采用AdamW优化器，初始学习率为0.001，并使用T_max=10的CosineAnnealingLR调度器；训练目标为Huber损失。批大小设为16，最大训练轮次为100，验证损失连续20轮未改善时触发早停，随机种子固定为1。训练启用自动混合精度。",
    )
    body_paragraph(
        doc,
        chapter5,
        "每个样本包含连续12个历史时间步，并同时预测未来12个时间步。图像输入由历史序列经Morlet连续小波变换生成，尺度范围为1—128；小波系数以flatten模式构造单通道128×128时频图。完整模型由ConvNeXt图像分支与STBAN时序分支组成，两分支输出通过可学习的全局Softmax权重进行融合；No-Fuse变体则以固定等权方式融合。主要设置汇总于表2。",
    )
    caption_paragraph(doc, chapter5, "表2  实验环境与主要参数设置", is_table=True)
    table2 = insert_table_clone(doc, chapter5, table_templates[1])
    write_table2(table2)

    heading_paragraph(doc, chapter5, "4.5  对比实验结果与分析", 2)
    heading_paragraph(doc, chapter5, "4.5.1  整体性能对比", 3)
    body_paragraph(
        doc,
        chapter5,
        "表3给出了9种模型在三个数据集上的12步平均误差。表中粗体表示同一数据集、同一指标下的最优值，下划线表示次优值。图2从数据集和指标两个维度展示原始误差，图3进一步给出各模型的相对排名，图4则直接比较CWT-TSI与每项指标上的最优基线。三类图从绝对误差、排序和相对变化三个角度构成对表3的互补说明。",
    )
    caption_paragraph(doc, chapter5, "表3  各模型在三个数据集上的整体性能对比", is_table=True)
    table3 = insert_table_clone(doc, chapter5, table_templates[2])
    write_comparison_table(table3, evidence["comparison"])

    add_figure(
        doc,
        chapter5,
        COMPARISON_FIGURE_DIR / "comparison_metric_grid.png",
        width_cm=15.6,
    )
    caption_paragraph(doc, chapter5, "图2  各模型在三个数据集上的整体误差对比")
    body_paragraph(
        doc,
        chapter5,
        "由图2可见，基线模型在数据集规模增大后表现出不同程度的误差上升，其中LSTM、ASTGCN-r、STSGCN和STFGNN在BJMetro上的绝对误差较高。CWT-TSI在HZMetro和BJMetro的三项指标上均取得最低值，在XMBRT上虽然MAE和RMSE不是最低，但MAPE低于全部基线。该结果表明，模型优势并非在所有数据集和所有指标上完全一致，需要结合排名和相对变化进一步分析。",
    )
    add_figure(
        doc,
        chapter5,
        COMPARISON_FIGURE_DIR / "comparison_rank_heatmap.png",
        width_cm=15.6,
    )
    caption_paragraph(doc, chapter5, "图3  各模型在数据集—指标组合上的排名")
    body_paragraph(
        doc,
        chapter5,
        "图3显示，CWT-TSI在9个“数据集—指标”组合中有7项排名第1，另外两项分别排名第2和第3，平均排名为1.33；GMAN的平均排名为2.00。与仅报告平均值相比，排名热图更清楚地揭示了CWT-TSI在中、大规模数据集上的一致性，同时也保留了其在XMBRT绝对误差指标上的不足。",
    )
    add_figure(
        doc,
        chapter5,
        COMPARISON_FIGURE_DIR / "comparison_cwt_vs_best_baseline.png",
        width_cm=15.6,
    )
    caption_paragraph(doc, chapter5, "图4  CWT-TSI相对最优基线的性能变化")
    body_paragraph(
        doc,
        chapter5,
        "图4以最优基线为参照，正值表示CWT-TSI误差更低，负值表示CWT-TSI误差更高。在XMBRT上，CWT-TSI的MAE和RMSE分别比最优基线高0.6%和1.3%，但MAPE降低6.5%；在HZMetro上三项指标分别降低1.2%、0.1%和9.6%；在BJMetro上分别降低5.4%、3.2%和7.5%。因此，CWT-TSI的主要优势集中在HZMetro、BJMetro以及三个数据集的相对误差控制，而不是对所有绝对误差指标的无条件支配。",
    )
    body_paragraph(
        doc,
        chapter5,
        "具体而言，XMBRT上GMAN取得最低MAE（5.3326），Graph WaveNet取得最低RMSE（8.6264），CWT-TSI的MAE为5.3655、RMSE为8.7420，分别排名第2和第3；其MAPE为35.7134，优于最强基线GMAN的38.1923。三项指标之间的差异说明，在该小规模数据集上，CWT时频表征对比例误差的改善较为明确，但未同步转化为最低的绝对误差。",
    )
    body_paragraph(
        doc,
        chapter5,
        "HZMetro上，CWT-TSI的MAE、RMSE和MAPE分别为10.7766、17.8384和24.2111，三项指标均排名第1。其MAE和RMSE相对GMAN仅改善1.2%和0.1%，二者在绝对误差上较为接近；MAPE则相对Graph WaveNet改善9.6%。这表明CWT-TSI在该数据集上的优势更多体现在相对误差，而RMSE差异较小。",
    )
    body_paragraph(
        doc,
        chapter5,
        "BJMetro上，CWT-TSI的MAE、RMSE和MAPE分别为16.8717、29.8812和18.6053，均为所有模型中的最低值；相对GMAN分别改善5.4%、3.2%和7.5%。在节点数最多的场景中，CWT-TSI仍保持三项指标一致领先，说明双分支模型在该实验设置下对大规模网络具有较好的适用性。由于本文未进行多随机种子重复试验或显著性检验，上述结论限定于已记录的一次完整实验结果，不作统计显著性推断。",
    )

    heading_paragraph(doc, chapter5, "4.5.2  各预测步长性能分析", 3)
    body_paragraph(
        doc,
        chapter5,
        "为考察预测距离增加时的误差变化，表4—表6列出CWT-TSI在三个数据集上第1至第12个预测步长的测试指标。所有数值均来自与表3中CWT-TSI相同的最终检查点，其逐步结果的总体汇总与表3对应。",
    )
    for table_no, dataset, template_idx in [
        (4, "XMBRT", 3),
        (5, "HZMetro", 4),
        (6, "BJMetro", 5),
    ]:
        caption_paragraph(
            doc,
            chapter5,
            f"表{table_no}  CWT-TSI在{dataset}数据集上的逐步预测性能",
            is_table=True,
        )
        step_table = insert_table_clone(doc, chapter5, table_templates[template_idx])
        write_step_table(step_table, evidence["full_runs"][dataset])
    add_figure(doc, chapter5, FIGURE_DIR / "cwt_per_step_metrics.png", width_cm=15.6)
    caption_paragraph(doc, chapter5, "图5  CWT-TSI在不同预测步长下的误差变化")
    body_paragraph(
        doc,
        chapter5,
        "图5表明，三个数据集的远期预测误差总体高于近端预测误差，但局部变化并非严格单调。XMBRT的最低MAE与RMSE均出现在第2步，最低MAPE出现在第5步；第12步相对各自最低值分别增加4.9%、6.0%和3.4%。其增长幅度小于另外两个数据集，说明该数据集上的12步预测误差相对平稳。",
    )
    body_paragraph(
        doc,
        chapter5,
        "HZMetro的MAE和RMSE由第1步的10.3356和16.9357总体上升至第12步的11.3211和18.9174，增幅分别为9.5%和11.7%；MAPE在第2步最低（23.5601%），第12步增至25.4943%，相对增幅为8.2%。除个别相邻步的小幅波动外，三项指标随预测距离增加呈较清晰的上升趋势。",
    )
    body_paragraph(
        doc,
        chapter5,
        "BJMetro的最低MAE出现在第2步，最低RMSE出现在第1步，最低MAPE出现在第4步。到第12步，三项指标相对各自最低值分别增加10.1%、10.9%和8.0%。这说明在节点规模最大的场景中，预测距离增加带来的误差累积更为明显；但第12步的三项结果仍包含在表3所示整体平均水平附近的连续变化范围内，没有出现突发式失稳。",
    )

    heading_paragraph(doc, chapter5, "4.6  消融实验分析", 2)
    body_paragraph(
        doc,
        chapter5,
        "为区分两个分支及融合方式的作用，本文构造三种变体：No-Img移除ConvNeXt图像分支，仅保留STBAN时序分支；No-TS移除时序分支，仅保留图像分支；No-Fuse保留双分支，但以固定等权平均替代可学习Softmax加权。各变体与完整模型使用相同的数据划分和主要训练设置。结果见表7，原始误差与相对变化分别见图6和图7。",
    )
    caption_paragraph(doc, chapter5, "表7  CWT-TSI及消融变体在三个数据集上的性能对比", is_table=True)
    table7 = insert_table_clone(doc, chapter5, table_templates[6])
    write_ablation_table(table7, evidence["ablation"])
    add_figure(
        doc,
        chapter5,
        ABLATION_FIGURE_DIR / "ablation_metric_grid.png",
        width_cm=15.6,
    )
    caption_paragraph(doc, chapter5, "图6  CWT-TSI及其消融变体的整体误差对比")
    body_paragraph(
        doc,
        chapter5,
        "图6显示，不同消融操作的影响具有明显的数据集依赖性。No-TS在HZMetro和BJMetro上产生最大的误差上升，而在XMBRT上与完整模型较为接近；No-Img的变化整体较小，并在HZMetro的MAPE以及BJMetro的MAE、MAPE上略优于完整模型；No-Fuse则在所有数据集和指标上均弱于完整模型。",
    )
    add_figure(
        doc,
        chapter5,
        ABLATION_FIGURE_DIR / "ablation_relative_change_heatmap.png",
        width_cm=15.6,
    )
    caption_paragraph(doc, chapter5, "图7  消融变体相对完整模型的误差变化")
    body_paragraph(
        doc,
        chapter5,
        "图7中的正值表示消融变体误差高于完整模型，负值表示变体误差更低。移除时序分支后，XMBRT三项指标仅增加0.1%—1.8%，但HZMetro分别增加25.6%、36.5%和8.6%，BJMetro分别增加29.7%、45.2%和18.2%。这一结果表明，时序分支对中、大规模数据集的绝对误差控制尤为重要，而小规模XMBRT中图像分支单独使用仍能保留大部分预测能力。",
    )
    body_paragraph(
        doc,
        chapter5,
        "移除图像分支后，XMBRT的MAE、RMSE和MAPE分别增加0.7%、0.5%和2.3%，说明时频图像在该数据集上提供了一定补充信息。HZMetro中No-Img的MAE和RMSE分别增加0.3%和0.4%，但MAPE降低0.1%；BJMetro中No-Img的RMSE增加0.3%，而MAE和MAPE分别降低约0.1%和1.4%。因此，图像分支的增益并非在所有场景和指标上一致，不能据此宣称完整模型在每一项消融比较中均最优。",
    )
    body_paragraph(
        doc,
        chapter5,
        "与固定等权融合相比，可学习Softmax融合在9项比较中均取得更低误差：No-Fuse在XMBRT、HZMetro和BJMetro上的误差增幅范围分别为0.6%—2.5%、0.6%—1.7%和1.3%—1.8%。这一结果支持“学习融合权重优于固定等权”的设计选择，但现有实验未保存样本级或时段级权重轨迹，因此只能验证全局可学习融合的有效性，不能进一步推断模型在特定高峰或扰动时段如何动态调整分支权重。",
    )
    body_paragraph(
        doc,
        chapter5,
        "综合来看，完整模型在9项指标中取得6项最优，No-Img取得其余3项最优；完整模型相对No-TS在全部9项指标上更优，相对No-Fuse也在全部9项指标上更优。消融实验由此确认了时序分支和可学习融合的稳定贡献，同时揭示出图像分支贡献随数据集规模和评价指标变化的边界条件。",
    )

    heading_paragraph(doc, chapter5, "4.7  模型收敛性分析", 2)
    body_paragraph(
        doc,
        chapter5,
        "图8给出三个数据集上完整CWT-TSI的训练损失和验证损失。曲线直接读取最终运行保存的history.npz，与表3、表4—表6所对应的模型检查点来自同一组实验。为避免将检查点选择与收敛趋势混为一谈，图中仅呈现损失轨迹，不附加检查点标记。",
    )
    add_figure(doc, chapter5, FIGURE_DIR / "cwt_training_convergence.png", width_cm=15.6)
    caption_paragraph(doc, chapter5, "图8  CWT-TSI在三个数据集上的训练与验证损失")
    body_paragraph(
        doc,
        chapter5,
        "三个数据集的训练与验证损失均在训练初期快速下降，随后进入下降速度减缓并伴有波动的平台阶段，最终由早停条件结束训练。XMBRT和HZMetro较早进入平台区，后期验证损失围绕较低水平波动；BJMetro的初始损失更高、优化历程更长，反映出其客流量级和节点规模带来的更大拟合难度。",
    )
    body_paragraph(
        doc,
        chapter5,
        "训练损失与验证损失在后期并非同步单调下降，说明继续增加训练轮次不能保证验证性能持续改善，采用验证集检查点和早停机制是必要的。需要指出的是，单条训练曲线只能用于描述本次运行的优化过程，不能替代多随机种子稳定性分析，也不能仅凭曲线形态断言不存在过拟合。总体而言，现有日志未出现损失发散，三个数据集均形成了可用于测试评估的稳定检查点。",
    )

    doc.save(str(OUTPUT_DOCX))
    print(f"docx=CREATED path={OUTPUT_DOCX}")
    return OUTPUT_DOCX


def table_text(table: Table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def audit_docx(path: Path, evidence: dict) -> dict:
    doc = Document(str(path))
    paragraphs = doc.paragraphs
    start = next(
        idx
        for idx, paragraph in enumerate(paragraphs)
        if paragraph.text.strip().startswith("4 实验结果与分析")
    )
    end = next(
        idx
        for idx, paragraph in enumerate(paragraphs)
        if paragraph.text.strip().startswith("5 结论与展望")
    )
    chapter_text = "\n".join(p.text for p in paragraphs[start:end])

    expected_headings = [
        "4.1  实验数据集",
        "4.1.1  XMBRT数据集",
        "4.1.2  HZMetro数据集",
        "4.1.3  BJMetro数据集",
        "4.2  对比基线模型",
        "4.3  评价指标",
        "4.4  实验环境与参数设置",
        "4.5  对比实验结果与分析",
        "4.5.1  整体性能对比",
        "4.5.2  各预测步长性能分析",
        "4.6  消融实验分析",
        "4.7  模型收敛性分析",
    ]
    structure_pass = all(heading in chapter_text for heading in expected_headings)
    if len(doc.tables) != 7:
        structure_pass = False

    data_pass = True
    comparison_table = doc.tables[2]
    for row_idx, model in enumerate(MODELS, start=2):
        col = 1
        for dataset in DATASETS:
            for metric in METRICS:
                actual = float(comparison_table.cell(row_idx, col).text)
                expected = evidence["comparison"][dataset][model][metric]
                if not math.isclose(actual, expected, abs_tol=5e-5):
                    data_pass = False
                col += 1
    for table_idx, dataset in zip([3, 4, 5], DATASETS):
        table = doc.tables[table_idx]
        for row_idx, step_row in enumerate(
            evidence["full_runs"][dataset]["per_step_metrics"], start=1
        ):
            for col_idx, metric in enumerate(METRICS, start=1):
                actual = float(table.cell(row_idx, col_idx).text)
                if not math.isclose(actual, step_row[metric], abs_tol=5e-5):
                    data_pass = False
        for col_idx, metric in enumerate(METRICS, start=1):
            actual = float(table.cell(13, col_idx).text)
            expected = evidence["full_runs"][dataset]["metrics"][metric]
            if not math.isclose(actual, expected, abs_tol=5e-5):
                data_pass = False
    ablation_table = doc.tables[6]
    row_idx = 1
    for dataset in DATASETS:
        for variant in VARIANTS:
            for col_idx, metric in enumerate(METRICS, start=2):
                actual = float(ablation_table.cell(row_idx, col_idx).text)
                expected = evidence["ablation"][dataset][variant][metric]
                if not math.isclose(actual, expected, abs_tol=5e-5):
                    data_pass = False
            row_idx += 1

    forbidden_patterns = [
        r"第90轮",
        r"第95轮",
        r"第94轮",
        r"最佳轮次",
        r"自动提升图像分支权重",
        r"自动提高图像分支权重",
        r"显著优于",
        r"MAE=5\.38、RMSE=8\.77、MAPE=36\.43",
    ]
    forbidden_hits = [
        pattern for pattern in forbidden_patterns if re.search(pattern, chapter_text)
    ]
    forbidden_pass = not forbidden_hits

    caption_numbers = [
        int(match)
        for match in re.findall(r"图([2-8])\s", chapter_text)
    ]
    figure_pass = len(doc.inline_shapes) >= 7 and set(caption_numbers) == set(range(2, 9))

    result = {
        "document": str(path),
        "structure_pass": structure_pass,
        "data_pass": data_pass,
        "figure_pass": figure_pass,
        "forbidden_claims_pass": forbidden_pass,
        "forbidden_hits": forbidden_hits,
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "chapter4_paragraph_count": end - start,
    }
    (EVIDENCE_DIR / "docx_audit.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if not all([structure_pass, data_pass, figure_pass, forbidden_pass]):
        raise AssertionError(json.dumps(result, ensure_ascii=False))
    print("data=PASS structure=PASS figures=PASS forbidden_claims=PASS")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-only", action="store_true")
    parser.add_argument("--figures-only", action="store_true")
    parser.add_argument("--build-docx", action="store_true")
    parser.add_argument("--audit-docx", type=Path)
    args = parser.parse_args()

    evidence = build_evidence()
    if args.audit_only:
        return
    if args.audit_docx:
        audit_docx(args.audit_docx.resolve(), evidence)
        return
    generate_figures(evidence)
    if args.figures_only:
        return
    output = build_docx(evidence)
    audit_docx(output, evidence)


if __name__ == "__main__":
    main()
